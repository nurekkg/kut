"""Позиции версии (ADR-015 часть 4, S6a-1): видимость позиций и цен по ролям, право на деньги, гарды
черновика, импорт XLSX. Гарды проверяются прямым SQL под kut_app без RETURNING - мимо сервиса."""

import io
from collections.abc import AsyncIterator
from decimal import Decimal
from typing import Any
from uuid import uuid4

import httpx
import openpyxl
import pytest
from httpx import AsyncClient
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from sqlalchemy import text
from sqlalchemy.engine import Engine
from sqlalchemy.exc import DBAPIError

from app.core.context import RequestContext, use_context
from app.core.db import get_portal_sessionmaker, get_sessionmaker
from app.core.redis import get_redis
from app.modules.auth import devseed
from tests.auth_support import AuthWorld
from tests.documents_support import (
    API,
    PDF,
    PDF_OK,
    XLSX,
    Desk,
    approve,
    create_document,
    leave_no_access_behind,  # noqa: F401
    open_desk,
    upload,
    zip_bytes,
)
from tests.equipment_support import idem, in_context, rows, second_project

HEADER = ["№", "Наименование", "Ед.", "Количество", "Цена"]


@pytest.fixture
async def s3() -> AsyncIterator[httpx.AsyncClient]:
    async with httpx.AsyncClient(timeout=10) as client:
        yield client


@pytest.fixture(autouse=True)
async def forget_the_granted_rights() -> AsyncIterator[None]:
    yield
    await get_redis().flushdb()


def xlsx(table: list[list[Any]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in table:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def line(name: str, quantity: str = "2", unit: str = "м3") -> dict[str, Any]:
    return {
        "name": name,
        "unit": unit,
        "quantity": quantity,
        "work_id": None,
        "project_node_id": None,
    }


def lines_url(document: str, version: str) -> str:
    return f"{API}/documents/{document}/versions/{version}/lines"


def ctx(auth_world: AuthWorld, who: str, desk: Desk) -> RequestContext:
    return in_context(auth_world.users[who], (desk.project,), (desk.legal,))


async def ks2_version(
    api: AsyncClient, s3: httpx.AsyncClient, desk: Desk, auth_world: AuthWorld
) -> tuple[str, str]:
    created = await create_document(
        api,
        desk.heads["estimator"],
        {
            "kind": "ks2",
            "legal_entity_id": str(desk.legal),
            "project_id": str(desk.project),
            "counterparty_id": str(auth_world.counterparties[devseed.CUSTOMER]),
            "title": "КС-2 с позициями",
        },
    )
    assert created.status_code == 201, created.text
    document = created.json()["id"]
    uploaded = await upload(api, s3, document, desk.heads["estimator"], PDF_OK, PDF)
    assert uploaded.status_code == 201, uploaded.text
    return document, uploaded.json()["id"]


async def put_lines(
    api: AsyncClient,
    headers: dict[str, str],
    document: str,
    version: str,
    lines: list[dict[str, Any]],
) -> httpx.Response:
    return await api.put(
        lines_url(document, version), json={"lines": lines}, headers={**headers, **idem()}
    )


async def put_prices(
    api: AsyncClient,
    headers: dict[str, str],
    document: str,
    version: str,
    prices: list[tuple[int, str]],
) -> httpx.Response:
    body = {"prices": [{"position_no": no, "unit_price": price} for no, price in prices]}
    return await api.put(
        f"{API}/documents/{document}/versions/{version}/line-prices",
        json=body,
        headers={**headers, **idem()},
    )


async def import_lines(
    api: AsyncClient,
    headers: dict[str, str],
    document: str,
    version: str,
    content: bytes,
    key: str | None = None,
) -> httpx.Response:
    keyed = {**headers, "Idempotency-Key": key} if key else {**headers, **idem()}
    return await api.post(f"{lines_url(document, version)}/import", content=content, headers=keyed)


async def write(sql: str, context: RequestContext, params: dict[str, Any]) -> None:
    """Запись мимо сервиса: SQL от kut_app в контексте пользователя, без RETURNING."""
    with use_context(context):
        async with get_sessionmaker()() as session:
            await session.execute(text(sql), params)
            await session.commit()


async def priced_draft(
    api: AsyncClient,
    s3: httpx.AsyncClient,
    desk: Desk,
    auth_world: AuthWorld,
) -> tuple[str, str]:
    document, version = await ks2_version(api, s3, desk, auth_world)
    lines_put = await put_lines(
        api,
        desk.heads["estimator"],
        document,
        version,
        [line("Бетон B25"), line("Арматура", "1.5", "т")],
    )
    assert lines_put.status_code == 200, lines_put.text
    priced = await put_prices(
        api, desk.heads["estimator"], document, version, [(1, "100.00"), (2, "50.50")]
    )
    assert priced.status_code == 200, priced.text
    return document, version


async def test_the_foreman_sees_the_positions_and_no_price_the_estimator_sees_both(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "foreman")
    document, version = await priced_draft(api, s3, desk, auth_world)

    estimator = await api.get(lines_url(document, version), headers=desk.heads["estimator"])
    assert estimator.status_code == 200, estimator.text
    assert Decimal(estimator.json()["total"]) == Decimal("275.75")
    assert Decimal(estimator.json()["lines"][0]["unit_price"]) == Decimal("100")

    foreman = await api.get(lines_url(document, version), headers=desk.heads["foreman"])
    assert foreman.status_code == 200, foreman.text
    body = foreman.json()
    assert [item["name"] for item in body["lines"]] == ["Бетон B25", "Арматура"]
    assert [Decimal(item["quantity"]) for item in body["lines"]] == [Decimal("2"), Decimal("1.5")]
    assert all(item["unit_price"] is None and item["amount"] is None for item in body["lines"])
    assert body["total"] is None
    assert "100.00" not in foreman.text and "50.50" not in foreman.text

    stored = await rows(
        "SELECT count(*) FROM document_line_price WHERE version_id = CAST(:v AS uuid)",
        ctx(auth_world, "foreman", desk),
        {"v": version},
    )
    assert stored == [(0,)]


async def test_a_new_list_keeps_the_prices_of_the_same_positions(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await priced_draft(api, s3, desk, auth_world)

    replaced = await put_lines(
        api,
        desk.heads["estimator"],
        document,
        version,
        [
            {"position_no": 1, **line("Бетон B30")},
            {"position_no": 2, **line("Арматура", "3", "т")},
            line("Щебень"),
        ],
    )
    assert replaced.status_code == 200, replaced.text
    items = replaced.json()["lines"]
    assert [item["name"] for item in items] == ["Бетон B30", "Арматура", "Щебень"]
    assert [item["unit_price"] for item in items[:2]] == ["100.00", "50.50"]
    assert items[2]["unit_price"] is None
    assert replaced.json()["total"] is None


async def test_a_list_with_prices_is_replaced_only_with_the_money_right(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(
        api, auth_world, worker_sync_engine, "estimator", "procurement", "foreman"
    )
    document, version = await priced_draft(api, s3, desk, auth_world)
    other, other_version = await ks2_version(api, s3, desk, auth_world)

    # Снабженец пишет позиции, но денег не видит: переписать список с ценами он не может.
    refused = await put_lines(
        api, desk.heads["procurement"], document, version, [line("Бетон новый")]
    )
    assert refused.status_code == 403, refused.text
    refused_import = await import_lines(
        api,
        desk.heads["procurement"],
        document,
        version,
        xlsx([HEADER, [1, "Бетон", "м3", 1, None]]),
    )
    assert refused_import.status_code == 403, refused_import.text
    # Пустой черновик без цен он вносит.
    accepted = await put_lines(
        api, desk.heads["procurement"], other, other_version, [line("Щебень")]
    )
    assert accepted.status_code == 200, accepted.text

    # Прораб не вносит позиции (нет documents:write:lines) и не ставит цены (нет права на деньги).
    assert (
        await put_lines(api, desk.heads["foreman"], document, version, [line("Бетон")])
    ).status_code == 403
    assert (
        await put_prices(api, desk.heads["foreman"], document, version, [(1, "1.00")])
    ).status_code == 403

    kept = await api.get(lines_url(document, version), headers=desk.heads["estimator"])
    assert [item["name"] for item in kept.json()["lines"]] == ["Бетон B25", "Арматура"]
    assert Decimal(kept.json()["lines"][0]["unit_price"]) == Decimal("100")


async def test_the_customer_sees_positions_and_prices_of_the_approved_version_only(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "customer", "sub_a")
    document, version = await priced_draft(api, s3, desk, auth_world)

    draft = await api.get(lines_url(document, version), headers=desk.heads["customer"])
    assert draft.status_code == 404, draft.text

    approve(worker_sync_engine, version)
    seen = await api.get(lines_url(document, version), headers=desk.heads["customer"])
    assert seen.status_code == 200, seen.text
    assert [item["name"] for item in seen.json()["lines"]] == ["Бетон B25", "Арматура"]
    assert [Decimal(item["unit_price"]) for item in seen.json()["lines"]] == [
        Decimal("100"),
        Decimal("50.50"),
    ]

    other = await api.get(lines_url(document, version), headers=desk.heads["sub_a"])
    assert other.status_code == 404, other.text


async def test_the_api_refuses_every_write_once_the_version_has_left_the_draft(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await priced_draft(api, s3, desk, auth_world)
    approve(worker_sync_engine, version)
    headers = desk.heads["estimator"]

    changed = await put_lines(api, headers, document, version, [line("Новая")])
    assert changed.status_code == 409, changed.text
    assert changed.json()["detail"] == "document.lines_not_draft"
    assert (await put_prices(api, headers, document, version, [(1, "5")])).status_code == 409
    imported = await import_lines(
        api, headers, document, version, xlsx([HEADER, [1, "А", "м", 1, None]])
    )
    assert imported.status_code == 409, imported.text

    kept = await api.get(lines_url(document, version), headers=headers)
    assert [item["name"] for item in kept.json()["lines"]] == ["Бетон B25", "Арматура"]


async def test_the_database_refuses_direct_writes_to_lines_and_prices_outside_the_draft(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await priced_draft(api, s3, desk, auth_world)
    approve(worker_sync_engine, version)
    context = ctx(auth_world, "estimator", desk)
    params = {"d": document, "v": version, "le": str(desk.legal), "p": str(desk.project)}

    with pytest.raises(DBAPIError, match=r"line\.version_not_draft"):
        await write(
            "INSERT INTO document_line (document_id, version_id, legal_entity_id, project_id,"
            " position_no, name, unit, quantity) VALUES (CAST(:d AS uuid), CAST(:v AS uuid),"
            " CAST(:le AS uuid), CAST(:p AS uuid), 9, 'Прямая', 'шт', 1)",
            context,
            params,
        )
    with pytest.raises(DBAPIError, match=r"line\.version_not_draft"):
        await write(
            "UPDATE document_line SET deleted_at = now() WHERE version_id = CAST(:v AS uuid)",
            context,
            params,
        )
    with pytest.raises(DBAPIError, match=r"line\.version_not_draft"):
        await write(
            "UPDATE document_line_price SET unit_price = 1 WHERE version_id = CAST(:v AS uuid)",
            context,
            params,
        )
    with pytest.raises(DBAPIError, match=r"line\.version_not_draft"):
        await write(
            "INSERT INTO document_line_price (line_id, document_id, version_id, legal_entity_id,"
            " project_id, unit_price) SELECT l.id, l.document_id, l.version_id, l.legal_entity_id,"
            " l.project_id, 7 FROM document_line l WHERE l.version_id = CAST(:v AS uuid)",
            context,
            params,
        )
    with pytest.raises(DBAPIError, match="permission denied"):
        await write(
            "DELETE FROM document_line WHERE version_id = CAST(:v AS uuid)", context, params
        )

    intact = await rows(
        "SELECT count(*), min(p.unit_price) FROM document_line l"
        " JOIN document_line_price p ON p.line_id = l.id"
        " WHERE l.version_id = CAST(:v AS uuid) AND l.deleted_at IS NULL",
        context,
        {"v": version},
    )
    assert intact == [(2, Decimal("50.50"))]


async def test_the_database_refuses_rows_that_point_at_another_document_version(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    first, first_version = await priced_draft(api, s3, desk, auth_world)
    second, second_version = await ks2_version(api, s3, desk, auth_world)
    context = ctx(auth_world, "estimator", desk)
    line_id = (
        await rows(
            "SELECT id FROM document_line WHERE version_id = CAST(:v AS uuid) ORDER BY position_no",
            context,
            {"v": first_version},
        )
    )[0][0]

    with pytest.raises(DBAPIError, match=r"document\.not_found"):
        await write(
            "INSERT INTO document_line (document_id, version_id, legal_entity_id, project_id,"
            " position_no, name, unit, quantity) VALUES (CAST(:d AS uuid), CAST(:v AS uuid),"
            " CAST(:le AS uuid), CAST(:p AS uuid), 9, 'Чужая', 'шт', 1)",
            context,
            {"d": first, "v": second_version, "le": str(desk.legal), "p": str(desk.project)},
        )
    with pytest.raises(DBAPIError, match=r"document\.not_found"):
        await write(
            "INSERT INTO document_line_price (line_id, document_id, version_id, legal_entity_id,"
            " project_id, unit_price) VALUES (CAST(:l AS uuid), CAST(:d AS uuid),"
            " CAST(:v AS uuid), CAST(:le AS uuid), CAST(:p AS uuid), 1)",
            context,
            {
                "l": str(line_id),
                "d": second,
                "v": second_version,
                "le": str(desk.legal),
                "p": str(desk.project),
            },
        )


async def test_a_writer_without_the_lines_right_cannot_add_a_line_and_cannot_change_a_price(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "foreman")
    document, version = await priced_draft(api, s3, desk, auth_world)
    foreman = ctx(auth_world, "foreman", desk)
    params = {"d": document, "v": version, "le": str(desk.legal), "p": str(desk.project)}

    with pytest.raises(DBAPIError, match=r"document\.forbidden|row-level security"):
        await write(
            "INSERT INTO document_line (document_id, version_id, legal_entity_id, project_id,"
            " position_no, name, unit, quantity) VALUES (CAST(:d AS uuid), CAST(:v AS uuid),"
            " CAST(:le AS uuid), CAST(:p AS uuid), 9, 'Прорабская', 'шт', 1)",
            foreman,
            params,
        )
    # Цены прораб не видит: политика отбирает строки, UPDATE их не находит и ничего не меняет.
    await write(
        "UPDATE document_line_price SET unit_price = 1 WHERE version_id = CAST(:v AS uuid)",
        foreman,
        params,
    )
    kept = await rows(
        "SELECT min(p.unit_price) FROM document_line_price p WHERE p.version_id = CAST(:v AS uuid)",
        ctx(auth_world, "estimator", desk),
        {"v": version},
    )
    assert kept == [(Decimal("50.50"),)]


async def test_the_import_takes_a_valid_file_and_refuses_a_bad_one_whole(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    headers = desk.heads["estimator"]

    good = xlsx([HEADER, [1, "Бетон B25", "м3", 2.5, 100], [2, "Арматура", "т", 1, None]])
    imported = await import_lines(api, headers, document, version, good)
    assert imported.status_code == 200, imported.text
    items = imported.json()["lines"]
    assert [item["name"] for item in items] == ["Бетон B25", "Арматура"]
    assert Decimal(items[0]["unit_price"]) == Decimal("100")
    assert items[1]["unit_price"] is None
    assert imported.json()["total"] is None

    bad = xlsx([HEADER, [1, "", "м3", 1, None], [2, "Ок", "м", -1, None], [3, "Ок", "м", 1, "abc"]])
    refused = await import_lines(api, headers, document, version, bad)
    assert refused.status_code == 422, refused.text
    detail = refused.json()["detail"]
    assert detail["code"] == "document.lines_invalid"
    assert [(item["row"], item["code"]) for item in detail["issues"]] == [
        (2, "name"),
        (3, "quantity"),
        (4, "price"),
    ]

    kept = await api.get(lines_url(document, version), headers=headers)
    assert [item["name"] for item in kept.json()["lines"]] == ["Бетон B25", "Арматура"]

    wrong_header = xlsx(
        [["Номер", "Наименование", "Ед.", "Количество", "Цена"], [1, "А", "м", 1, None]]
    )
    assert (await import_lines(api, headers, document, version, wrong_header)).json()["detail"] == {
        "code": "document.lines_header"
    }
    empty = xlsx([HEADER])
    assert (await import_lines(api, headers, document, version, empty)).json()["detail"] == {
        "code": "document.lines_empty"
    }
    not_a_book = await import_lines(api, headers, document, version, b"not a workbook")
    assert not_a_book.json()["detail"] == {"code": "document.lines_format"}


async def test_the_import_refuses_formulas_macros_and_too_many_or_too_big_files(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    headers = desk.heads["estimator"]

    formula = xlsx([HEADER, [1, "Бетон", "м3", 2, "=D2*2"]])
    refused = await import_lines(api, headers, document, version, formula)
    assert refused.status_code == 422, refused.text
    assert refused.json()["detail"] == {"code": "document.lines_formula", "rows": [2]}

    macros = zip_bytes(
        {
            "[Content_Types].xml": b"<Types>sheet.macroEnabled.main+xml</Types>",
            "xl/workbook.xml": b"<workbook/>",
            "xl/vbaProject.bin": b"\x00",
        }
    )
    refused = await import_lines(api, headers, document, version, macros)
    assert refused.json()["detail"] == {"code": "document.lines_macros"}

    too_many = xlsx([HEADER] + [[n, f"Позиция {n}", "шт", 1, None] for n in range(1, 2002)])
    refused = await import_lines(api, headers, document, version, too_many)
    assert refused.json()["detail"] == {"code": "document.lines_too_many"}

    too_big = await import_lines(api, headers, document, version, b"x" * (2 * 1024 * 1024 + 1))
    assert too_big.status_code == 413, too_big.text
    assert too_big.json()["detail"] == "document.lines_too_large"

    assert (
        await rows(
            "SELECT count(*) FROM document_line WHERE version_id = CAST(:v AS uuid)",
            ctx(auth_world, "estimator", desk),
            {"v": version},
        )
    ) == [(0,)]


async def test_a_repeated_import_with_the_same_key_answers_the_same_and_writes_once(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    headers = desk.heads["estimator"]
    content = xlsx([HEADER, [1, "Бетон", "м3", 2, None], [2, "Арматура", "т", 1, None]])
    key = str(uuid4())

    first = await import_lines(api, headers, document, version, content, key=key)
    again = await import_lines(api, headers, document, version, content, key=key)
    assert first.status_code == again.status_code == 200, again.text
    assert again.json() == first.json()
    stored = await rows(
        "SELECT count(*) FROM document_line WHERE version_id = CAST(:v AS uuid)",
        ctx(auth_world, "estimator", desk),
        {"v": version},
    )
    assert stored == [(2,)]


async def test_the_template_is_an_xlsx_with_the_header_row(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    response = await api.get(
        f"{API}/documents/lines/template.xlsx", headers=desk.heads["estimator"]
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert [cell.value for cell in sheet[1]] == HEADER


async def test_a_kind_without_positions_refuses_them(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    created = await create_document(
        api,
        desk.heads["estimator"],
        {
            "kind": "commissioning_act",
            "legal_entity_id": str(desk.legal),
            "project_id": str(desk.project),
            "title": "Акт без позиций",
        },
    )
    assert created.status_code == 201, created.text
    uploaded = await upload(api, s3, created.json()["id"], desk.heads["estimator"], PDF_OK, PDF)
    assert uploaded.status_code == 201, uploaded.text

    refused = await put_lines(
        api,
        desk.heads["estimator"],
        created.json()["id"],
        uploaded.json()["id"],
        [line("Не для акта")],
    )
    assert refused.status_code == 422, refused.text
    assert refused.json()["detail"] == "document.lines_not_allowed"


async def test_the_audit_of_prices_is_visible_only_with_the_money_right(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "foreman")
    _, version = await priced_draft(api, s3, desk, auth_world)
    sql = (
        "SELECT count(*) FROM audit_log WHERE table_name = :t"
        " AND COALESCE(new, old) ->> 'version_id' = :v"
    )

    foreman_prices = await rows(
        sql, ctx(auth_world, "foreman", desk), {"t": "document_line_price", "v": version}
    )
    estimator_prices = await rows(
        sql, ctx(auth_world, "estimator", desk), {"t": "document_line_price", "v": version}
    )
    foreman_lines = await rows(
        sql, ctx(auth_world, "foreman", desk), {"t": "document_line", "v": version}
    )
    assert foreman_prices == [(0,)]
    assert estimator_prices[0][0] >= 2
    assert foreman_lines[0][0] >= 2


async def test_the_kinds_say_which_kinds_take_positions(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    response = await api.get(f"{API}/documents/kinds", headers=desk.heads["estimator"])
    assert response.status_code == 200, response.text
    flags = {item["code"]: item["lines"] for item in response.json()}
    assert flags["ks2"] is True
    assert flags["commissioning_act"] is False


async def test_the_portal_login_sees_the_lines_of_an_approved_version_only_in_sql(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "customer")
    document, version = await priced_draft(api, s3, desk, auth_world)
    portal = RequestContext(
        user_id=auth_world.users["customer"],
        is_internal=False,
        project_ids=frozenset({desk.project}),
        legal_entity_ids=frozenset(),
        counterparty_id=auth_world.counterparties[devseed.CUSTOMER],
        request_id=None,
        sees_all_users=False,
    )
    counts = (
        "SELECT (SELECT count(*) FROM document_line WHERE version_id = CAST(:v AS uuid)),"
        " (SELECT count(*) FROM document_line_price WHERE version_id = CAST(:v AS uuid))"
    )
    with use_context(portal):
        async with get_portal_sessionmaker()() as session:
            draft = (await session.execute(text(counts), {"v": version})).one()
    assert tuple(draft) == (0, 0), document

    approve(worker_sync_engine, version)
    with use_context(portal):
        async with get_portal_sessionmaker()() as session:
            approved = (await session.execute(text(counts), {"v": version})).one()
    assert tuple(approved) == (2, 2)


async def test_deleting_a_position_keeps_the_prices_of_the_others(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    """Цена держится за номером позиции: удаление средней строки не сдвигает цены соседей."""
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    headers = desk.heads["estimator"]
    created = await put_lines(api, headers, document, version, [line("А"), line("Б"), line("В")])
    assert created.status_code == 200, created.text
    await put_prices(api, headers, document, version, [(1, "100.00"), (2, "200.00"), (3, "300.00")])

    replaced = await put_lines(
        api,
        headers,
        document,
        version,
        [{"position_no": 1, **line("А")}, {"position_no": 3, **line("В")}],
    )
    assert replaced.status_code == 200, replaced.text
    prices = {item["position_no"]: item["unit_price"] for item in replaced.json()["lines"]}
    assert prices == {1: "100.00", 3: "300.00"}


async def test_a_new_position_gets_a_new_number_and_no_deleted_price(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await priced_draft(api, s3, desk, auth_world)
    replaced = await put_lines(
        api,
        desk.heads["estimator"],
        document,
        version,
        [{"position_no": 1, **line("Бетон B25")}, line("Новая")],
    )
    assert replaced.status_code == 200, replaced.text
    items = [(item["position_no"], item["unit_price"]) for item in replaced.json()["lines"]]
    assert items == [(1, "100.00"), (3, None)]


async def test_a_duplicate_position_number_and_a_blank_name_are_refused(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    headers = desk.heads["estimator"]
    duplicate = await put_lines(
        api,
        headers,
        document,
        version,
        [{"position_no": 1, **line("А")}, {"position_no": 1, **line("Б")}],
    )
    assert duplicate.status_code == 422, duplicate.text
    assert duplicate.json()["detail"] == "document.lines_invalid"
    blank = await put_lines(api, headers, document, version, [line("   ")])
    assert blank.status_code == 422, blank.text


async def test_a_node_of_another_object_is_refused_by_the_database_guard(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    other = second_project(worker_sync_engine, desk.legal)
    with worker_sync_engine.begin() as connection:
        foreign_node = connection.scalar(
            text(
                "INSERT INTO project_node (project_id, legal_entity_id, type, name)"
                " VALUES (CAST(:p AS uuid), CAST(:le AS uuid), 'building', 'Чужой корпус')"
                " RETURNING id"
            ),
            {"p": str(other), "le": str(desk.legal)},
        )
    refused = await put_lines(
        api,
        desk.heads["estimator"],
        document,
        version,
        [{"position_no": 1, **line("Бетон"), "project_node_id": str(foreign_node)}],
    )
    assert refused.status_code == 422, refused.text
    assert refused.json()["detail"] == "document.line_invalid"
    with pytest.raises(DBAPIError, match=r"line\.node_invalid"):
        await write(
            "INSERT INTO document_line (document_id, version_id, legal_entity_id, project_id,"
            " position_no, name, unit, quantity, project_node_id) VALUES (CAST(:d AS uuid),"
            " CAST(:v AS uuid), CAST(:le AS uuid), CAST(:p AS uuid), 9, 'Чужая', 'шт', 1,"
            " CAST(:n AS uuid))",
            ctx(auth_world, "estimator", desk),
            {
                "d": document,
                "v": version,
                "le": str(desk.legal),
                "p": str(desk.project),
                "n": str(foreign_node),
            },
        )


async def test_an_array_formula_and_a_broken_workbook_are_refused_not_stored(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator")
    document, version = await ks2_version(api, s3, desk, auth_world)
    headers = desk.heads["estimator"]

    book = Workbook()
    sheet = book.active
    sheet.append(HEADER)
    sheet.append([1, "Бетон", "м3", ArrayFormula("D2", "=2*2"), None])
    stream = io.BytesIO()
    book.save(stream)
    array = await import_lines(api, headers, document, version, stream.getvalue())
    assert array.status_code == 422, array.text
    assert array.json()["detail"] == {"code": "document.lines_formula", "rows": [2]}

    broken = zip_bytes({"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<workbook"})
    refused = await import_lines(api, headers, document, version, broken)
    assert refused.status_code == 422, refused.text
    assert refused.json()["detail"] == {"code": "document.lines_format"}


async def test_a_money_user_without_the_lines_right_cannot_change_prices(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    """Бухгалтер видит деньги, но позиции не вносит: сменить существующую цену он не может.
    Отказ - 403, а не сбой сервера."""
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "accountant")
    document, version = await priced_draft(api, s3, desk, auth_world)
    refused = await put_prices(api, desk.heads["accountant"], document, version, [(1, "999.00")])
    assert refused.status_code == 403, refused.text
    assert refused.json()["detail"] == "forbidden"
