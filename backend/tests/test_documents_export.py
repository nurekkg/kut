"""Выгрузка реестра в XLSX (S5b): право documents:export, суммы только с documents:read:financial,
формулы в текстовых ячейках экранированы, строки и фильтры - как у реестра."""

from collections.abc import AsyncIterator
from io import BytesIO
from typing import Any

import httpx
import pytest
from httpx import AsyncClient, Response
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import text
from sqlalchemy.engine import Engine

from app.core.redis import get_redis
from app.modules.auth import devseed
from tests.auth_support import AuthWorld
from tests.documents_support import (
    PDF,
    PDF_OK,
    Desk,
    approve,
    create_document,
    leave_no_access_behind,  # noqa: F401
    open_desk,
    upload,
)

API = "/api/v1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
async def s3() -> AsyncIterator[httpx.AsyncClient]:
    async with httpx.AsyncClient(timeout=10) as client:
        yield client


async def make(
    api: AsyncClient,
    desk: Desk,
    who: str,
    body: dict[str, Any],
    s3: httpx.AsyncClient,
    engine: Engine,
    *,
    approved: bool = False,
) -> dict[str, Any]:
    response = await create_document(
        api,
        desk.heads[who],
        {
            "legal_entity_id": str(desk.legal),
            "project_id": str(desk.project),
            "title": "Выгрузка",
            **body,
        },
    )
    assert response.status_code == 201, response.text
    document = response.json()
    uploaded = await upload(api, s3, document["id"], desk.heads[who], PDF_OK, PDF)
    assert uploaded.status_code == 201, uploaded.text
    if approved:
        approve(engine, uploaded.json()["id"])
    return document


async def export(api: AsyncClient, headers: dict[str, str], **params: str) -> Response:
    return await api.get(f"{API}/documents/export.xlsx", params=params, headers=headers)


def sheet_of(response: Response) -> Worksheet:
    return load_workbook(BytesIO(response.content)).active


def rows_of(response: Response) -> tuple[tuple[Any, ...], list[tuple[Any, ...]]]:
    header, *rows = sheet_of(response).iter_rows(values_only=True)
    return header, rows


async def test_the_foreman_and_the_subcontractor_get_no_export(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "foreman", "sub_a")
    for who in ("foreman", "sub_a"):
        response = await export(api, desk.heads[who], project_id=str(desk.project))
        assert response.status_code == 403, response.text


async def test_the_accountant_downloads_the_registry_with_its_sums_as_an_attachment(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "accountant")
    customer = auth_world.counterparties[devseed.CUSTOMER]
    ks2 = await make(
        api,
        desk,
        "estimator",
        {"kind": "ks2", "counterparty_id": str(customer), "amount": "98000.00"},
        s3,
        worker_sync_engine,
        approved=True,
    )

    response = await export(api, desk.heads["accountant"], project_id=str(desk.project))

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == 'attachment; filename="documents.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    header, rows = rows_of(response)
    assert header[-1] == "Сумма"
    row = next(item for item in rows if item[0] == ks2["registration_no"])
    assert row[-1] == 98000.0
    assert row[5] == "Утверждён"


async def test_the_foreman_with_a_temporary_export_right_gets_a_file_without_sums(
    api: AsyncClient,
    auth_world: AuthWorld,
    worker_sync_engine: Engine,
    owner_sync_engine: Engine,
    s3: httpx.AsyncClient,
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "foreman")
    customer = auth_world.counterparties[devseed.CUSTOMER]
    ks2 = await make(
        api,
        desk,
        "estimator",
        {"kind": "ks2", "counterparty_id": str(customer), "amount": "98000.00"},
        s3,
        worker_sync_engine,
        approved=True,
    )
    with owner_sync_engine.begin() as connection:
        grant = connection.execute(
            text(
                "INSERT INTO role_permission (role_id, permission_id)"
                " SELECT :role, p.id FROM permission p"
                " WHERE p.module = 'documents' AND p.action = 'export' AND p.qualifier IS NULL"
                " RETURNING id"
            ),
            {"role": auth_world.roles["foreman"]},
        ).scalar_one()
    try:
        await get_redis().flushdb()  # кэш прав живёт до минуты: сбрасываем, чтобы не ждать
        response = await export(api, desk.heads["foreman"], project_id=str(desk.project))
        assert response.status_code == 200, response.text
        header, rows = rows_of(response)
        assert "Сумма" not in header
        assert any(item[0] == ks2["registration_no"] for item in rows)
        assert all(98000.0 not in item for item in rows)
    finally:
        with owner_sync_engine.begin() as connection:
            connection.execute(
                text("UPDATE role_permission SET deleted_at = now() WHERE id = :id"),
                {"id": grant},
            )
        await get_redis().flushdb()


async def test_formulas_in_text_cells_are_escaped_and_no_cell_is_a_formula(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "accountant")
    formula = await make(
        api,
        desk,
        "estimator",
        {"kind": "estimate_local", "title": '=HYPERLINK("http://example.com")'},
        s3,
        worker_sync_engine,
    )
    mention = await make(
        api,
        desk,
        "estimator",
        {"kind": "estimate_local", "title": "@cmd"},
        s3,
        worker_sync_engine,
    )

    response = await export(api, desk.heads["accountant"], project_id=str(desk.project))

    assert response.status_code == 200, response.text
    _, rows = rows_of(response)
    titles = {item[0]: item[2] for item in rows}
    assert titles[formula["registration_no"]] == "'" + formula["title"]
    assert titles[mention["registration_no"]] == "'@cmd"
    cells = [cell for row in sheet_of(response).iter_rows() for cell in row]
    assert all(cell.data_type != "f" for cell in cells)


async def test_the_export_takes_the_registry_filters(
    api: AsyncClient, auth_world: AuthWorld, worker_sync_engine: Engine, s3: httpx.AsyncClient
) -> None:
    desk = await open_desk(api, auth_world, worker_sync_engine, "estimator", "accountant")
    approved = await make(
        api, desk, "estimator", {"kind": "estimate_local"}, s3, worker_sync_engine, approved=True
    )
    draft = await make(api, desk, "estimator", {"kind": "estimate_local"}, s3, worker_sync_engine)

    response = await export(
        api, desk.heads["accountant"], project_id=str(desk.project), status="draft"
    )

    assert response.status_code == 200, response.text
    _, rows = rows_of(response)
    numbers = [item[0] for item in rows]
    assert draft["registration_no"] in numbers
    assert approved["registration_no"] not in numbers
